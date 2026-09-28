"""
Independent verification pass.

Part 1 re-derives every headline number from the raw per-seed files WITHOUT reusing
viva_reanalysis.py, so an error in that script cannot hide itself.
Part 2 checks that the number actually written into the workbook, the slides and the
response document matches.
"""
import ast
import json
import pickle
import re
import warnings
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

warnings.filterwarnings("ignore")

FYP = Path("/sessions/fervent-epic-einstein/mnt/FYP")
M = FYP / "FYP code" / "metrics"
VP = FYP / "viva_presentation"
XLSX = VP / "FYP_Viva_Result_with_Viva_Responses.xlsx"
PPTX = VP / "FYP_Viva_Presentation_Slide_REVISED.pptx"
DOCX = VP / "Viva_Response_OngShiKai_2026-08-05.docx"

fails, checks = [], 0


def check(name, ok, detail=""):
    global checks
    checks += 1
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'   ' + detail if detail else ''}")
    if not ok:
        fails.append(name)


def near(a, b, tol=5e-3):
    return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))


# ─────────────────────────────────────────── Part 1: re-derive from raw
print("\n=== PART 1 — independent re-derivation from the per-seed source files ===")
mh = pd.read_csv(M / "multi_horizon_results.csv")
raw = {}
for _, r in mh.iterrows():
    h = int(r["horizon"])
    g = lambda c: np.asarray(ast.literal_eval(
        str(r[c]).replace("np.float64(", "").replace(")", "")), dtype=float)
    raw[h] = dict(ilt=g("ilt_raw_runs"), lstm=g("lstm_raw_runs"), rf=g("rf_raw_runs"),
                  sarima=float(r["sarima_rmse"]))

H = [1, 7, 14, 21, 28]
check("30 seeds present at every horizon", all(len(raw[h]["ilt"]) == 30 for h in H))

# H1 — pooled reduction and the >=20% margin
ilt_all = np.concatenate([raw[h]["ilt"] for h in H])
sar_all = np.concatenate([np.full(30, raw[h]["sarima"]) for h in H])
red = 100 * (1 - ilt_all.mean() / sar_all.mean())
check("H1 pooled reduction = 47.5%", near(red, 47.534, 1e-3), f"got {red:.3f}%")
per = [100 * (1 - raw[h]["ilt"].mean() / raw[h]["sarima"]) for h in H]
check("H1 per-horizon range 42.5-52.0%", near(min(per), 42.475) and near(max(per), 52.020),
      f"got {min(per):.2f}-{max(per):.2f}%")
check("H1 30/30 wins at every horizon",
      all((raw[h]["ilt"] < raw[h]["sarima"]).sum() == 30 for h in H))
worst_p = max(
    stats.ttest_1samp(raw[h]["ilt"], 0.8 * raw[h]["sarima"]).pvalue / 2 for h in H)
check("H1 margin test: every Holm p < 1e-23", worst_p < 1e-23, f"worst raw p = {worst_p:.2e}")

# H2 — one-sided p at h=1 and pooled
d1 = raw[1]["ilt"] - raw[1]["lstm"]
t1, p1_two = stats.ttest_rel(raw[1]["ilt"], raw[1]["lstm"])
check("H2 h=1d two-sided p = 0.019 (the pre-viva figure)", near(p1_two, 0.0190, 0.02),
      f"got {p1_two:.4f}")
check("H2 h=1d one-sided p = 0.0095", near(p1_two / 2, 0.00949, 0.01), f"got {p1_two/2:.5f}")
praw = []
for h in H:
    t, p2 = stats.ttest_rel(raw[h]["ilt"], raw[h]["lstm"])
    praw.append(p2 / 2 if t < 0 else 1 - p2 / 2)
holm = []
order = np.argsort(praw)
run = 0.0
adj = np.empty(5)
for rank, idx in enumerate(order):
    run = max(run, (5 - rank) * praw[idx])
    adj[idx] = min(run, 1.0)
check("H2 h=1d Holm-corrected p = 0.047", near(adj[0], 0.0475, 0.02), f"got {adj[0]:.4f}")
dp = np.concatenate([raw[h]["ilt"] - raw[h]["lstm"] for h in H])
tp, p2p = stats.ttest_rel(np.concatenate([raw[h]["ilt"] for h in H]),
                          np.concatenate([raw[h]["lstm"] for h in H]))
check("H2 pooled one-sided p = 0.0056", near(p2p / 2, 0.005551, 0.02), f"got {p2p/2:.6f}")
check("H2 pooled mean diff = -0.537", near(dp.mean(), -0.5366, 1e-2), f"got {dp.mean():.4f}")
check("H2 pooled dz = -0.21", near(dp.mean() / dp.std(ddof=1), -0.20997, 1e-2),
      f"got {dp.mean()/dp.std(ddof=1):.4f}")
check("H2 mean RMSE 33.13 vs 33.67",
      near(ilt_all.mean(), 33.129, 1e-3) and
      near(np.concatenate([raw[h]["lstm"] for h in H]).mean(), 33.669, 1e-3),
      f"got {ilt_all.mean():.3f} vs {np.concatenate([raw[h]['lstm'] for h in H]).mean():.3f}")

# H3 — binomial
dv = pd.read_csv(M / "ch5_final_h3_domain_validation_RECONSTRUCTED.csv")
bt = stats.binomtest(int(dv.Domain_validated.sum()), len(dv), 0.70, alternative="greater")
ci = bt.proportion_ci(confidence_level=0.95, method="exact")
check("H3 exact binomial p = 7.98e-04", near(bt.pvalue, 7.9792e-4, 1e-3), f"got {bt.pvalue:.3e}")
check("H3 p equals 0.70^20 analytically", near(bt.pvalue, 0.70 ** 20, 1e-9))
check("H3 CI lower bound = 0.861", near(ci.low, 0.86089, 1e-3), f"got {ci.low:.5f}")
check("H3 CI lower bound exceeds the 0.70 criterion", ci.low > 0.70)
cons = pickle.load(open(M / "h3_30seed_attention_consistency.pkl", "rb"))
check("H3 cross-seed RMSE 31.93 +/- 1.95",
      near(cons["rmse_mean"], 31.929, 1e-3) and near(cons["rmse_std"], 1.9453, 1e-2))

# H4 — ratios and scaling
eff = pd.read_csv(M / "efficiency_table.csv").set_index("Model")
L, I = eff.loc["Standalone LSTM"], eff.loc["ILT (proposed)"]
check("H4 params ratio = 1.23", near(I.Params / L.Params, 1.2266, 1e-3))
check("H4 size ratio = 1.23", near(I["Model size (KB)"] / L["Model size (KB)"], 1.2337, 1e-3))
check("H4 latency ratio = 1.58",
      near(I["Inference per sample (ms)"] / L["Inference per sample (ms)"], 1.5817, 1e-3))
check("H4 train-time ratio = 1.77",
      near(I["Train time (s)"] / L["Train time (s)"], 1.7680, 1e-3))
se = I["Train time std (s)"] / np.sqrt(30)
t = (I["Train time (s)"] - 2 * L["Train time (s)"]) / se
p_tt = stats.t.cdf(t, 29)
check("H4 train-time p = 0.140 (fail to reject)", near(p_tt, 0.1399, 2e-2) and p_tt > 0.05,
      f"got {p_tt:.4f}")
hw = stats.t.ppf(0.975, 29) * se
lo_r = (I["Train time (s)"] - hw) / L["Train time (s)"]
hi_r = (I["Train time (s)"] + hw) / L["Train time (s)"]
check("H4 train-time ratio CI [1.34, 2.20] contains 2",
      near(lo_r, 1.337, 1e-2) and near(hi_r, 2.199, 1e-2) and lo_r < 2 < hi_r,
      f"got [{lo_r:.3f}, {hi_r:.3f}]")
se_l = I["Inference per sample std (ms)"] / np.sqrt(30)
p_lat = stats.t.cdf((I["Inference per sample (ms)"] - 2 * L["Inference per sample (ms)"]) / se_l, 29)
check("H4 latency p = 0.041 (reject)", near(p_lat, 0.0407, 3e-2) and p_lat < 0.05,
      f"got {p_lat:.4f}")
sc = pd.read_csv(M / "bigo_scaling_experiment.csv")
lr = stats.linregress(np.log(sc["T"]), np.log(sc["TemporalAttention_ms"]))
lr2 = stats.linregress(np.log(sc["T"]), np.log(sc["MultiHeadAttention_ms"]))
check("H4 TemporalAttention exponent = 0.53", near(lr.slope, 0.5259, 1e-2), f"got {lr.slope:.4f}")
check("H4 self-attention exponent = 1.42", near(lr2.slope, 1.4177, 1e-2), f"got {lr2.slope:.4f}")
tc = stats.t.ppf(0.975, len(sc) - 2)
check("H4 exponent CI [0.31, 0.74]",
      near(lr.slope - tc * lr.stderr, 0.3118, 1e-2) and near(lr.slope + tc * lr.stderr, 0.7399, 1e-2))
ratio_1792 = sc.MultiHeadAttention_ms.iloc[-1] / sc.TemporalAttention_ms.iloc[-1]
check("H4 49x faster at T = 1792", near(ratio_1792, 49.0, 1e-2), f"got {ratio_1792:.1f}x")

# H5 — bootstrap reproducibility with an independent seed
fc = pickle.load(open(M / "ilt_holdout_forecasts.pkl", "rb"))


def r2_of(h):
    p_, a_ = np.asarray(fc[h]["pred"], float).ravel(), np.asarray(fc[h]["actual"], float).ravel()
    return 1 - np.sum((a_ - p_) ** 2) / np.sum((a_ - a_.mean()) ** 2)


check("H5 h=1d ensemble R2 = 0.602", near(r2_of(1), 0.6016, 1e-3), f"got {r2_of(1):.4f}")
check("H5 h=7d ensemble R2 = 0.477", near(r2_of(7), 0.4767, 1e-3), f"got {r2_of(7):.4f}")
check("H5 h=14d R2 positive but h=21/28d negative",
      r2_of(14) > 0 > r2_of(21) > r2_of(28))
# rerun the bootstrap under a DIFFERENT seed — the decision must not depend on it
rng = np.random.default_rng(20260804)
stable = True
for h, expect_reject in [(1, True), (7, True), (14, False), (21, False), (28, False)]:
    p_, a_ = np.asarray(fc[h]["pred"], float).ravel(), np.asarray(fc[h]["actual"], float).ravel()
    n, B, BL = len(a_), 2000, 28
    starts = np.arange(0, n - BL + 1)
    boot = np.empty(B)
    for b in range(B):
        idx = np.concatenate([np.arange(s, s + BL) for s in rng.choice(starts, int(np.ceil(n / BL)))])[:n]
        pp, aa = p_[idx], a_[idx]
        boot[b] = 1 - np.sum((aa - pp) ** 2) / np.sum((aa - aa.mean()) ** 2)
    p_boot = max((boot <= 0).mean(), 1 / B)
    if (p_boot < 0.05) != expect_reject:
        stable = False
        print(f"      h={h}d decision flipped under a new bootstrap seed (p = {p_boot:.4f})")
check("H5 bootstrap decisions stable under a different RNG seed", stable)

# Model comparison mean ranks
ranks = []
for h in H:
    mat = np.vstack([np.full(30, raw[h]["sarima"]), raw[h]["rf"], raw[h]["lstm"], raw[h]["ilt"]]).T
    ranks.append(stats.rankdata(mat, axis=1))
mr = np.vstack(ranks).mean(axis=0)
check("Mean ranks ILT 1.42 / LSTM 1.58 / SARIMA 3.00 / RF 4.00",
      near(mr[3], 1.42, 1e-2) and near(mr[2], 1.58, 1e-2) and mr[0] == 3.0 and mr[1] == 4.0,
      f"got {np.round(mr, 3).tolist()}")
check("Ranks sum to 10 per cell (sanity)", near(mr.sum(), 10.0, 1e-9))

# RF overfitting
rf = pd.read_csv(M / "rf_feature_parity_multihorizon.csv")
tr_col = [c for c in rf.columns if "train" in c.lower() and "r2" in c.lower()][0]
te_col = [c for c in rf.columns if "test" in c.lower() and "r2" in c.lower()][0]
check("RF train R2 ~0.96 and test R2 in [-0.38, -0.30]",
      near(rf[tr_col].mean(), 0.962, 1e-2) and -0.38 < rf[te_col].mean() < -0.30,
      f"train {rf[tr_col].mean():.3f}, test {rf[te_col].mean():.3f}")

# ─────────────────────────────────────────── Part 1b: nb17 / nb18 / nb19
print("\n=== PART 1b — re-derivation from the new per-seed outputs ===")

ps = pd.read_csv(M / "nb17_finetune_arms_per_seed.csv")
check("nb17 · 900 rows = 6 arms x 5 horizons x 30 seeds", len(ps) == 900, f"got {len(ps)}")
a4 = ps[ps.arm == "A4_recalibrate"].groupby("horizon").r2.mean().reindex(H)
a0 = ps[ps.arm == "A0_frozen"].groupby("horizon").r2.mean().reindex(H)
check("nb17 A4 R2 = 0.69/0.56/0.45/0.30/0.22",
      all(near(a, b, 2e-2) for a, b in zip(a4.values, [0.6901, 0.5597, 0.4452, 0.3031, 0.2230])),
      f"got {np.round(a4.values, 4).tolist()}")
check("nb17 A4 positive at every horizon", (a4 > 0).all())
check("nb17 A4 degrades monotonically with horizon (the profile of real skill)",
      all(a4.values[i] > a4.values[i + 1] for i in range(4)))
check("nb17 A0 negative at h=14/21/28d and positive at h=1/7d",
      a0[1] > 0 and a0[7] > 0 and a0[14] < 0 and a0[21] < 0 and a0[28] < 0,
      f"got {np.round(a0.values, 4).tolist()}")
sr = ps[ps.arm == "A0_frozen"].groupby("horizon").scale_ratio.mean().reindex(H)
check("nb17 A0 over-disperses, ratio rising 1.24 -> 1.74",
      near(sr[1], 1.240, 2e-2) and near(sr[28], 1.741, 2e-2) and all(np.diff(sr.values) > 0),
      f"got {np.round(sr.values, 3).tolist()}")
a3 = ps[ps.arm == "A3_rescale"].groupby("horizon").r2.mean().reindex(H)
check("nb17 A3 R2 is near-flat across horizons (the caveat we state)",
      (a3.max() - a3.min()) < 0.05, f"spread {a3.max()-a3.min():.4f}")
a5 = ps[ps.arm == "A5_retrain_expanding"].groupby("horizon").r2.mean().reindex(H)
check("nb17 A5 retraining is worse than A0 at every horizon", (a5 < a0).all())

h3h = pd.read_csv(M / "nb18_h3_horizon_tests.csv")
check("nb18 H3 20/20 at all five horizons", (h3h.consensus_rate == 1.0).all() and len(h3h) == 5)
check("nb18 H3 rejects H0 at every horizon", (h3h.decision == "Reject H0").all())
check("nb18 H3 Holm p = 0.004 at every horizon", all(near(x, 0.004, 5e-2) for x in h3h.p_holm))

eff18 = pd.read_csv(M / "nb18_h4_efficiency_per_horizon_per_seed.csv")
piv = eff18.pivot_table(index=["horizon", "seed"], columns="model", values="train_s")
ratio = piv["ILT (proposed)"] / piv["Standalone LSTM"]
check("nb18 H4 paired train-time ratios span 0.83-1.35 by horizon",
      near(ratio.groupby("horizon").mean().min(), 0.8279, 2e-2)
      and near(ratio.groupby("horizon").mean().max(), 1.3521, 2e-2),
      f"got {ratio.groupby('horizon').mean().round(3).tolist()}")
check("nb18 H4 pairing shrinks the ILT SD from 14.4 s to ~3.9 s",
      near(eff18[eff18.model == "ILT (proposed)"].groupby("horizon").train_s.std().mean(), 3.93, 8e-2),
      f"got {eff18[eff18.model=='ILT (proposed)'].groupby('horizon').train_s.std().mean():.2f} s")
h4h = pd.read_csv(M / "nb18_h4_horizon_tests.csv")
check("nb18 H4 rejects H0 on every stochastic metric at every horizon",
      (h4h[h4h.p_holm.notna()].decision == "Reject H0").all())

ws = pd.read_csv(M / "nb18_window_sweep_tests.csv")
best = ws.loc[ws.groupby("horizon").rmse_mean.idxmin()]
check("nb18 window sweep: 7 d beats 28 d at both horizons",
      (best.seq_len == 7).all(), f"best seq_len per horizon {best.seq_len.tolist()}")
check("nb18 window sweep: 7 d significant after Holm at both horizons",
      (ws[ws.seq_len == 7].vs_28_p_holm < 0.05).all(),
      f"got {ws[ws.seq_len==7].vs_28_p_holm.round(4).tolist()}")

cv = pd.read_csv(M / "nb19_cv_design_summary.csv").set_index("strategy")
s3 = cv[cv.index.str.startswith("S3")].iloc[0]
s6 = cv[cv.index.str.startswith("S6")].iloc[0]
check("nb19 S6 produced folds this time (the bug that crashed the figure is fixed)", s6.n_fits > 0)
check("nb19 leakage: S6 RMSE 69% below S3",
      near(100 * (s3.rmse_mean - s6.rmse_mean) / s3.rmse_mean, 69.2, 2e-2),
      f"got {100*(s3.rmse_mean-s6.rmse_mean)/s3.rmse_mean:.1f}%")
check("nb19 leakage: S6 R2 = 0.97 vs S3 0.19",
      near(s6.r2_mean, 0.9702, 2e-2) and near(s3.r2_mean, 0.1855, 5e-2))
sr19 = pd.read_csv(M / "nb19_split_ratio_tests.csv")
spread = sr19.rmse_mean.max() - sr19.rmse_mean.min()
check("nb19 split-ratio spread = 22.2 RMSE", near(spread, 22.23, 2e-2), f"got {spread:.2f}")
check("nb19 spread is ~41x the ILT-vs-LSTM effect of 0.54",
      near(spread / 0.5366, 41.4, 5e-2), f"got {spread/0.5366:.1f}x")
check("nb19 split-ratio ordering is non-monotonic in training size (period, not volume)",
      not (sr19.sort_values("train_frac").rmse_mean.is_monotonic_increasing
           or sr19.sort_values("train_frac").rmse_mean.is_monotonic_decreasing))

# ─────────────────────────────────────────── Part 2: what the deliverables say
print("\n=== PART 2 — do the deliverables actually contain these numbers? ===")


def sheet_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        out[ws.title] = " ".join(str(c.value) for row in ws.iter_rows() for c in row
                                 if c.value is not None)
    return out


sh = sheet_text(XLSX)
alltext = " ".join(sh.values())

# Spot-check that the recomputed numbers actually reached the workbook. Structure,
# RQ labelling and "originals untouched" are covered by revision/verify_addendum.py.
for label, needle, where in [
    ("H1 margin null stated", "0.80 × RMSE(SARIMA)", "H1 (Accuracy vs SARIMA)"),
    ("H1 47.5% reduction", "47.5", None),
    ("H2 Holm p = 0.047", "0.047", "H2 (vs Standalone LSTM)"),
    ("H2 pooled p = 0.0056", "0.0056", "H2 (vs Standalone LSTM)"),
    ("H3 binomial p = 7.98e-04", "7.98e-04", "H3 (Domain Validation)"),
    ("H3 CI lower 0.861", "0.861", "H3 (Domain Validation)"),
    ("H3 holds at all five horizons", "cases_diff_14", "H3 (Domain Validation)"),
    ("H4 single-split p retained on record", "0.1399", "H4 (Computational Efficiency)"),
    ("H4 per-horizon ratio 0.828", "0.828", "H4 (Computational Efficiency)"),
    ("H4 pairing explanation present", "3.9 s", "H4 (Computational Efficiency)"),
    ("H5 h=14d p = 0.158", "0.158", "H5 (Real-World Generalisation)"),
    ("H5 recalibration recovery", "0.6901", "H5 (Real-World Generalisation)"),
    ("H5 over-dispersion diagnosis", "1.74", "H5 (Real-World Generalisation)"),
    ("H5 A3 shrinkage caveat", "variance reduction", "H5 (Real-World Generalisation)"),
    ("H5 notebook-label correction", "correction to our own tooling", "H5 (Real-World Generalisation)"),
    ("mean rank 1.42", "1.42", "Model Comparison (added)"),
    ("Friedman omnibus", "1.6", "Model Comparison (added)"),
    ("consistency audit lists the H2 two-sided error", "TWO-sided", "Consistency Audit (added)"),
    ("window sweep admits 28 d is not the optimum", "not the empirical optimum", "Methodology"),
    ("leakage quantified", "69", "Methodology"),
    ("paired-design defence", "one fortieth", "Methodology"),
]:
    hay = sh[where] if where else alltext
    check(f"workbook · {label}", needle in hay)

check("workbook · original sheets all still present",
      all(n in sh for n in ["Read Me", "Contribution", "RQ-RO-Hypothesis Map", "Dataset Sources",
                            "Methodology", "Seeds (Reproducibility)", "H1 (Accuracy vs SARIMA)",
                            "H2 (vs Standalone LSTM)", "H3 (Domain Validation)",
                            "H4 (Computational Efficiency)", "H5 (Real-World Generalisation)"]))

def pptx_text(path):
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        return {n: re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8")) for n in names}


pt = pptx_text(PPTX)
allp = " ".join(pt.values())
check("slides · H₀ and Hₐ both present", "H₀" in allp and "Hₐ" in allp)
check("slides · every hypothesis has a null", allp.count("H₀") >= 10)
check("slides · both decision outcomes appear", "REJECT H₀" in allp and "FAIL TO REJECT" in allp)
check("slides · H3 p-value on the deck", "7.98e-04" in allp)
check("slides · H4 now shows the per-horizon result", "0.83" in allp and "1.35" in allp)
check("slides · slide 4 carries the per-horizon H4 claim", "paired within each seed" in allp)
check("slides · H5 recalibration result on the deck", "recalibration" in allp.lower())
check("slides · 7 slides", len(pt) == 7, f"got {len(pt)}")

with zipfile.ZipFile(DOCX) as z:
    dt = re.sub(r"<[^>]+>", " ", z.read("word/document.xml").decode("utf-8"))
for label, needle in [
    ("all nine items addressed", None),
    ("H1 margin correction described", "0.80 × RMSE(SARIMA)"),
    ("H2 two-sided/one-sided correction described", "0.0095"),
    ("H3 binomial p", "7.98"),
    ("H4 reversal stated", "Widened back to SUPPORTED"),
    ("H4 withdrawal still on record", "withdrawn"),
    ("H5 recovery stated", "0.69"),
    ("H5 caveat stated", "variance reduction"),
    ("window-sweep result admitted", "7-day input window beats"),
    ("leakage quantified", "69%"),
    ("six adaptation arms listed", "A5 retrain from scratch"),
    ("six splitting strategies listed", "S6"),
    ("notebook filenames given", "Q6_17_finetune_patch_experiment.ipynb"),
]:
    if needle:
        check(f"response doc · {label}", needle in dt)
check("response doc · all nine examiner headings present",
      all(w in dt for w in ["Hypothesis Formulation", "Statistical Results", "Consistency",
                            "Structural Flow", "Model Comparison", "Fine-tuning Analysis",
                            "Visualisation", "Temporal Analysis", "Data Strategy"]))

# notebooks exist and parse
import ast as _ast
for nb_name in ["Q6_17_finetune_patch_experiment.ipynb", "Q8_18_all_windows_coverage.ipynb",
                "Q9_19_dataset_splitting_analysis.ipynb"]:
    p = FYP / "FYP code" / nb_name
    ok = p.exists()
    if ok:
        nb = json.load(open(p, encoding="utf-8"))
        for c in nb["cells"]:
            if c["cell_type"] == "code":
                try:
                    _ast.parse("".join(l + "\n" for l in c["source"]))
                except SyntaxError:
                    ok = False
    check(f"notebook · {nb_name} exists and every code cell parses", ok)
check("notebook · viva_revision_common.py present",
      (FYP / "FYP code" / "viva_revision_common.py").exists())

print(f"\n{'='*70}\n{checks - len(fails)} of {checks} checks passed")
if fails:
    print("FAILURES:")
    for f in fails:
        print("  -", f)
else:
    print("No discrepancies found.")
